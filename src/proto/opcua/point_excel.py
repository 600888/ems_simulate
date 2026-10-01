"""Read and validate the four-sheet OPC UA point template without DB writes."""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
import math
from typing import Any

from asyncua import ua
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

SHEETS: dict[str, tuple[int, tuple[str, ...]]] = {
    "遥测": (
        0,
        (
            "测点编码",
            "测点名称",
            "属性编码",
            "NodeId",
            "命名空间URI",
            "数据类型",
            "采样间隔(ms)",
            "单位",
            "乘系数",
            "加系数",
            "上限值",
            "下限值",
            "初始值",
        ),
    ),
    "遥信": (
        1,
        ("测点编码", "测点名称", "属性编码", "NodeId", "命名空间URI", "数据类型", "采样间隔(ms)", "是否反转", "初始值"),
    ),
    "遥控": (
        2,
        (
            "测点编码",
            "测点名称",
            "属性编码",
            "NodeId",
            "命名空间URI",
            "数据类型",
            "采样间隔(ms)",
            "命令类型",
            "初始值",
            "关联遥信编码",
        ),
    ),
    "遥调": (
        3,
        (
            "测点编码",
            "测点名称",
            "属性编码",
            "NodeId",
            "命名空间URI",
            "数据类型",
            "采样间隔(ms)",
            "单位",
            "乘系数",
            "加系数",
            "上限值",
            "下限值",
            "初始值",
            "关联遥测编码",
        ),
    ),
}
MAX_EXCEL_BYTES = 5 * 1024 * 1024
MAX_ROWS_PER_SHEET = 10000


@dataclass(frozen=True)
class ExcelError:
    sheet: str
    row: int
    field: str
    message: str

    def to_dict(self) -> dict[str, Any]:
        return {"sheet": self.sheet, "row": self.row, "field": self.field, "message": self.message}


@dataclass
class ParsedPoints:
    sha256: str
    rows: list[dict[str, Any]]
    errors: list[ExcelError]


def _required_text(value: Any, max_length: int) -> str:
    result = str(value).strip() if value is not None else ""
    if not result or len(result) > max_length:
        raise ValueError(f"必须填写 1 到 {max_length} 个字符")
    return result


def _finite_number(value: Any) -> float:
    if isinstance(value, bool) or value is None:
        raise ValueError("必须是有限数值")
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("必须是有限数值") from exc
    if not math.isfinite(number):
        raise ValueError("必须是有限数值")
    return number


def _initial_value(value: Any, data_type: str) -> bool | int | float:
    if data_type == "Boolean":
        if isinstance(value, str) and value.strip().lower() in {"true", "false", "0", "1"}:
            return value.strip().lower() in {"true", "1"}
        if value in (True, False, 0, 1):
            return bool(value)
        raise ValueError("Boolean 初始值必须为 0、1、true 或 false")
    if data_type == "Int32":
        if isinstance(value, bool) or not isinstance(value, int) or not -(2**31) <= value < 2**31:
            raise ValueError("Int32 初始值必须为 32 位整数")
        return value
    return _finite_number(value)


def _parse_row(sheet: str, row_number: int, values: tuple[Any, ...], errors: list[ExcelError]) -> dict[str, Any]:
    headers = SHEETS[sheet][1]
    raw = dict(zip(headers, values, strict=False))
    result: dict[str, Any] = {"sheet": sheet, "excel_row": row_number, "point_type": SHEETS[sheet][0]}

    def check(field: str, fn, default=None):
        try:
            return fn(raw.get(field))
        except ValueError as exc:
            errors.append(ExcelError(sheet, row_number, field, str(exc)))
            return default

    result["point_code"] = check("测点编码", lambda v: _required_text(v, 128))
    result["point_name"] = check("测点名称", lambda v: _required_text(v, 255))
    result["attribute_code"] = check("属性编码", lambda v: _required_text(v, 255))
    result["node_id"] = check("NodeId", lambda v: _required_text(v, 512))
    result["namespace_index"] = None
    if result["node_id"]:
        try:
            node_id = ua.NodeId.from_string(result["node_id"])
            if (
                not 1 <= node_id.NamespaceIndex <= 65535
                or not isinstance(node_id.Identifier, str)
                or not node_id.Identifier
            ):
                raise ValueError("必须是带命名空间索引的字符串 NodeId，例如 ns=2;s=... ")
            result["node_id"] = node_id.to_string()
            result["namespace_index"] = node_id.NamespaceIndex
        except (ValueError, AttributeError, ua.UaStringParsingError) as exc:
            errors.append(ExcelError(sheet, row_number, "NodeId", str(exc)))
    result["namespace_uri"] = check("命名空间URI", lambda v: _required_text(v, 255))
    result["data_type"] = check("数据类型", lambda v: _required_text(v, 32))
    allowed = {"Boolean"} if result["point_type"] in (1, 2) else {"Int32", "Double"}
    if result["data_type"] and result["data_type"] not in allowed:
        errors.append(ExcelError(sheet, row_number, "数据类型", f"当前 Sheet 仅支持 {', '.join(sorted(allowed))}"))
    interval = check("采样间隔(ms)", lambda v: int(v) if isinstance(v, int) and not isinstance(v, bool) else -1)
    if interval is None or not 1 <= interval <= 3600000:
        errors.append(ExcelError(sheet, row_number, "采样间隔(ms)", "必须是 1 到 3600000 的整数"))
    result["sampling_interval_ms"] = interval
    result["initial_value"] = None
    if result["data_type"] in allowed:
        result["initial_value"] = check("初始值", lambda v: _initial_value(v, result["data_type"]))

    result.update(
        unit=None,
        scale_mul=None,
        scale_add=None,
        upper_limit=None,
        lower_limit=None,
        reverse=False,
        command_type=None,
        related_code=None,
    )
    if result["point_type"] in (0, 3):
        result["unit"] = str(raw.get("单位") or "").strip() or None
        if result["unit"] and len(result["unit"]) > 64:
            errors.append(ExcelError(sheet, row_number, "单位", "最多 64 个字符"))
        for field, key in (
            ("乘系数", "scale_mul"),
            ("加系数", "scale_add"),
            ("上限值", "upper_limit"),
            ("下限值", "lower_limit"),
        ):
            result[key] = check(field, _finite_number)
        upper, lower = result["upper_limit"], result["lower_limit"]
        if upper is not None and lower is not None and upper < lower:
            errors.append(ExcelError(sheet, row_number, "上限值", "上限值不能小于下限值"))
        initial = result["initial_value"]
        if initial is not None and upper is not None and lower is not None and not lower <= initial <= upper:
            errors.append(ExcelError(sheet, row_number, "初始值", "初始值不在上下限范围内"))
    elif result["point_type"] == 1:
        reverse = raw.get("是否反转")
        if reverse not in (0, 1, True, False):
            errors.append(ExcelError(sheet, row_number, "是否反转", "必须为 0 或 1"))
        result["reverse"] = bool(reverse)
    else:
        result["command_type"] = check("命令类型", lambda v: _required_text(v, 64))
    if result["point_type"] in (2, 3):
        field = "关联遥信编码" if result["point_type"] == 2 else "关联遥测编码"
        result["related_code"] = check(field, lambda v: _required_text(v, 128))
    return result


def parse_point_excel(content: bytes) -> ParsedPoints:
    if not content or len(content) > MAX_EXCEL_BYTES:
        raise ValueError("Excel 文件为空或超过 5 MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 .xlsx 点表文件") from exc
    rows: list[dict[str, Any]] = []
    errors: list[ExcelError] = []
    try:
        for sheet, (_, expected_headers) in SHEETS.items():
            if sheet not in workbook.sheetnames:
                errors.append(ExcelError(sheet, 1, "Sheet", "缺少必需 Sheet"))
                continue
            worksheet = workbook[sheet]
            iterator = worksheet.iter_rows(values_only=True)
            actual_headers = tuple(next(iterator, ()))[: len(expected_headers)]
            if actual_headers != expected_headers:
                errors.append(ExcelError(sheet, 1, "表头", "表头与 OPC UA 点表模板不一致"))
                continue
            for row_number, values in enumerate(iterator, start=2):
                if row_number > MAX_ROWS_PER_SHEET + 1:
                    errors.append(ExcelError(sheet, row_number, "Sheet", "每个 Sheet 最多 10000 行"))
                    break
                if all(value is None for value in values):
                    continue
                rows.append(_parse_row(sheet, row_number, values, errors))
        for extra in set(workbook.sheetnames) - SHEETS.keys():
            errors.append(ExcelError(extra, 1, "Sheet", "存在非模板 Sheet"))
    finally:
        workbook.close()

    for key, field in (("point_code", "测点编码"), ("node_id", "NodeId")):
        seen: dict[str, dict[str, Any]] = {}
        for row in rows:
            value = row.get(key)
            if not value:
                continue
            if value in seen:
                errors.append(
                    ExcelError(
                        row["sheet"],
                        row["excel_row"],
                        field,
                        f"与 {seen[value]['sheet']} 第 {seen[value]['excel_row']} 行重复",
                    )
                )
            else:
                seen[value] = row
    index_uris: dict[int, str] = {}
    uri_indexes: dict[str, int] = {}
    for row in rows:
        index, uri = row.get("namespace_index"), row.get("namespace_uri")
        if index is None or not uri:
            continue
        if index in index_uris and index_uris[index] != uri:
            errors.append(
                ExcelError(row["sheet"], row["excel_row"], "命名空间URI", f"命名空间索引 {index} 对应多个 URI")
            )
        if uri in uri_indexes and uri_indexes[uri] != index:
            errors.append(ExcelError(row["sheet"], row["excel_row"], "NodeId", "同一命名空间 URI 对应多个索引"))
        index_uris[index], uri_indexes[uri] = uri, index
    return ParsedPoints(sha256(content).hexdigest(), rows, errors)


def export_point_excel(points: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    common_fields = (
        "point_code",
        "point_name",
        "attribute_code",
        "node_id",
        "namespace_uri",
        "data_type",
        "sampling_interval_ms",
    )
    analog_fields = ("unit", "scale_mul", "scale_add", "upper_limit", "lower_limit", "initial_value")
    extra_fields = {
        0: analog_fields,
        1: ("reverse", "initial_value"),
        2: ("command_type", "initial_value", "related_code"),
        3: (*analog_fields, "related_code"),
    }
    try:
        for sheet, (point_type, headers) in SHEETS.items():
            worksheet = workbook.create_sheet(sheet)
            worksheet.append(headers)
            fields = (*common_fields, *extra_fields[point_type])
            for point in points:
                if point["point_type"] != point_type:
                    continue
                worksheet.append([point.get(field) for field in fields])
                for cell in worksheet[worksheet.max_row]:
                    if isinstance(cell.value, str):
                        # A point name or code is text even when it starts with '='.
                        cell.data_type = "s"
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2F5597")
                worksheet.column_dimensions[cell.column_letter].width = 22 if cell.column < 4 else 28
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()
    finally:
        workbook.close()
