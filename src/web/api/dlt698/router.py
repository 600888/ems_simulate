"""HTTP contracts for the thin dlt698 adapter."""

import asyncio
from collections import defaultdict
from functools import wraps
from io import BytesIO
import json

from dlt698.errors import Dlt698Error
from fastapi import APIRouter, HTTPException, Request, UploadFile
from fastapi.routing import APIRoute
from openpyxl import load_workbook

from src.data.service.channel_service import ChannelService
from src.data.service.dlt698_service import Dlt698Service
from src.device.protocol.dlt698_handler import Dlt698Handler
from src.device.protocol.runtime_config import normalize_protocol_params
from src.proto.dlt698.catalog import catalog, standard_definition
from src.proto.dlt698.schema import (
    AutoReadRequest,
    Dlt698Config,
    Operation,
    Point,
    ProbeRequest,
    ReadRequest,
    RecordQuery,
)
from src.web.api.channel.helpers import reload_device_instance
from src.web.api.exceptions import ConflictError, NotFoundError, ServiceUnavailableError, ValidationError
from src.web.api.schemas import BaseResponse


class Dlt698Route(APIRoute):
    def get_route_handler(self):
        handler = super().get_route_handler()

        @wraps(handler)
        async def checked(request):
            try:
                return await handler(request)
            except (Dlt698Error, ConnectionError, TimeoutError) as error:
                raise ServiceUnavailableError(str(error)) from error
            except (ValueError, TypeError, KeyError) as error:
                raise ValidationError(str(error)) from error

        return checked


router = APIRouter(prefix="/api/dlt698/channels/{channel_id}", tags=["dlt698"], route_class=Dlt698Route)
_locks = defaultdict(asyncio.Lock)


async def while_connected(request, operation):
    """客户端取消 HTTP 读取后，取消原生在途事务并释放通道锁。"""
    task = asyncio.create_task(operation)

    async def disconnected():
        # 请求体已由 FastAPI 解析完毕。持续等待 ASGI 断线消息，避免
        # is_disconnected() 的即时取消检查被 BaseHTTPMiddleware 的检查点吞掉。
        while (await request.receive())["type"] != "http.disconnect":
            pass

    disconnect = asyncio.create_task(disconnected())
    try:
        done, _ = await asyncio.wait({task, disconnect}, return_when=asyncio.FIRST_COMPLETED)
        if task in done:
            return task.result()
        await disconnect
        raise HTTPException(status_code=499, detail="客户端已取消读取")
    finally:
        for pending in (task, disconnect):
            if not pending.done():
                pending.cancel()
        await asyncio.gather(task, disconnect, return_exceptions=True)


async def get_handler(channel_id, request):
    channel = await asyncio.to_thread(ChannelService.get_channel_by_id, channel_id)
    if channel is None or channel.get("protocol_type") != 8:
        raise NotFoundError("DL/T 698 通道不存在")
    controller = request.app.state.device_controller
    device = controller.get_device_by_id(channel_id)
    if device is None:
        await reload_device_instance(controller, channel_id, is_start=False)
        device = controller.get_device_by_id(channel_id)
    if not device or not isinstance(device.protocol_handler, Dlt698Handler):
        raise ServiceUnavailableError("DL/T 698 设备实例未就绪")
    return device.protocol_handler


async def refresh(channel_id, request):
    controller = request.app.state.device_controller
    device = controller.get_device_by_id(channel_id)
    await reload_device_instance(controller, channel_id, is_start=bool(device and device.is_protocol_running()))


@router.get("/status")
async def status(channel_id: int, request: Request):
    return BaseResponse(data=(await get_handler(channel_id, request)).status())


@router.get("/config")
async def config(channel_id: int):
    return BaseResponse(data=(await asyncio.to_thread(Dlt698Service.config, channel_id)).model_dump())


@router.put("/config")
async def save_config(channel_id: int, body: Dlt698Config, request: Request):
    async with _locks[channel_id]:
        channel = await asyncio.to_thread(ChannelService.get_channel_by_id, channel_id)
        if not channel or channel.get("protocol_type") != 8:
            raise NotFoundError("DL/T 698 通道不存在")
        normalize_protocol_params(8, channel["conn_type"], body.model_dump())
        points = await asyncio.to_thread(Dlt698Service.points, channel_id)
        for point in points:
            standard_definition(Point.model_validate(point), body)
        if body.role == "server":
            from src.proto.dlt698.runtime import Dlt698Runtime

            await asyncio.to_thread(Dlt698Runtime(body, {}, points)._make_device)
        await asyncio.to_thread(Dlt698Service.save_config, channel_id, body)
        await refresh(channel_id, request)
        return BaseResponse(data=body.model_dump())


@router.get("/points")
async def points(channel_id: int, request: Request):
    runtime = (await get_handler(channel_id, request)).get_runtime()
    definitions = await asyncio.to_thread(Dlt698Service.points, channel_id)
    return BaseResponse(data=[{**point, **runtime.snapshots.get(point["code"], {})} for point in definitions])


async def persist_points(channel_id, body, request, original=None):
    async with _locks[channel_id]:
        handler = await get_handler(channel_id, request)
        persisted = await asyncio.to_thread(Dlt698Service.points, channel_id) if original else []
        previous = next((point for point in persisted if point["code"] == original), None)
        data = await asyncio.to_thread(Dlt698Service.save_points, channel_id, body, original)
        if (
            previous
            and len(data) == 1
            and all(previous[key] == data[0][key] for key in ("code", "oad", "data_type", "initial_value", "scaler"))
        ):
            # Metadata-only edits retain publication, history, mappings and simulation state.
            runtime = handler.get_runtime()
            point = runtime.point_states.get(original)
            if point is not None:
                definition = runtime.definitions[original]
                current_value = definition["initial_value"]
                definition.update(data[0])
                definition["initial_value"] = current_value
                point.name = data[0]["name"]
                point.min_value_limit = data[0]["min_value_limit"]
                point.max_value_limit = data[0]["max_value_limit"]
                return data
        await refresh(channel_id, request)
        return data


@router.post("/points/import")
async def import_points(channel_id: int, request: Request, file: UploadFile):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise ValidationError("请使用 .xlsx，initial_value 列填写有类型的 Data JSON")
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise ValidationError("点表文件不能超过 5 MiB")

    def parse():
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                raise ValueError("工作簿为空")
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, ())
            required = {"code", "name", "oad", "data_type", "initial_value"}
            if not required.issubset(set(header)):
                raise ValueError("点表缺少必需列: " + ", ".join(sorted(required)))
            definitions = []
            for number, values in enumerate(rows, start=2):
                if all(value is None for value in values):
                    continue
                if len(definitions) >= 5000:
                    raise ValueError("一次最多导入 5000 个测点")
                definition = {
                    str(key): value for key, value in zip(header, values, strict=False) if key and value is not None
                }
                try:
                    definition["initial_value"] = json.loads(definition["initial_value"])
                    definitions.append(Point.model_validate(definition))
                except (ValueError, TypeError) as error:
                    raise ValueError(f"第 {number} 行: {error}") from error
            if not definitions:
                raise ValueError("点表没有数据")
            return definitions
        finally:
            workbook.close()

    definitions = await asyncio.to_thread(parse)
    await persist_points(channel_id, definitions, request)
    await asyncio.to_thread(ChannelService.update_channel, channel_id, dlt698_point_mode="import")
    return BaseResponse(data={"count": len(definitions)})


@router.post("/points")
async def add_point(channel_id: int, body: Point, request: Request):
    return BaseResponse(data=(await persist_points(channel_id, [body], request))[0])


@router.post("/points/standard")
async def import_standard_points(channel_id: int, request: Request):
    async with _locks[channel_id]:
        await get_handler(channel_id, request)
        count = await asyncio.to_thread(Dlt698Service.import_standard_points, channel_id)
        await asyncio.to_thread(ChannelService.update_channel, channel_id, dlt698_point_mode="standard")
        await refresh(channel_id, request)
        return BaseResponse(data={"count": count})


@router.put("/points/{code}")
async def edit_point(channel_id: int, code: str, body: Point, request: Request):
    return BaseResponse(data=(await persist_points(channel_id, [body], request, code))[0])


@router.delete("/points/{code}")
async def delete_point(channel_id: int, code: str, request: Request):
    async with _locks[channel_id]:
        await get_handler(channel_id, request)
        await asyncio.to_thread(Dlt698Service.delete_point, channel_id, code)
        await refresh(channel_id, request)
        return BaseResponse(data=True)


@router.post("/catalog")
async def point_catalog(channel_id: int, body: dict):
    current = await asyncio.to_thread(Dlt698Service.config, channel_id)
    merged = Dlt698Config.model_validate({**current.model_dump(), **body})
    return BaseResponse(data=await asyncio.to_thread(catalog, merged))


@router.post("/read")
async def read(channel_id: int, body: ReadRequest, request: Request):
    async def operation():
        async with _locks[channel_id]:
            runtime = (await get_handler(channel_id, request)).get_runtime()
            if any(code not in runtime.definitions for code in body.codes):
                raise ValidationError("请求包含不存在的测点编码")
            return await runtime.read([runtime.definitions[code]["oad"] for code in body.codes], body.codes)

    return BaseResponse(data=await while_connected(request, operation()))


@router.post("/probe")
async def probe(channel_id: int, body: ProbeRequest, request: Request):
    async def operation():
        async with _locks[channel_id]:
            return await (await get_handler(channel_id, request)).get_runtime().probe(body.oads)

    return BaseResponse(data=await while_connected(request, operation()))


@router.post("/operations")
async def operations(channel_id: int, body: Operation, request: Request):
    async with _locks[channel_id]:
        return BaseResponse(data=await (await get_handler(channel_id, request)).get_runtime().operate(body))


@router.post("/records/query")
async def query_records(channel_id: int, body: RecordQuery, request: Request):
    async with _locks[channel_id]:
        return BaseResponse(data=await (await get_handler(channel_id, request)).get_runtime().query_records(body))


@router.post("/records/rows")
async def append_record(channel_id: int, request: Request):
    await get_handler(channel_id, request)
    raise ConflictError("dlt698 1.0 托管服务端不提供记录 provider 注入；客户端记录查询已接入")


@router.get("/reports")
async def reports(channel_id: int, request: Request):
    return BaseResponse(data=list((await get_handler(channel_id, request)).get_runtime().reports))


@router.get("/routes")
async def routes(channel_id: int, request: Request):
    await get_handler(channel_id, request)
    return BaseResponse(data=[])


@router.post("/routes")
async def save_route(channel_id: int, request: Request):
    await get_handler(channel_id, request)
    raise ConflictError("代理路由需要 dlt698 专家接口，当前托管适配器未启用")


@router.put("/auto-read")
async def auto_read(channel_id: int, body: AutoReadRequest, request: Request):
    async with _locks[channel_id]:
        return BaseResponse(
            data=await (await get_handler(channel_id, request))
            .get_runtime()
            .set_auto(body.running, body.interval_ms, body.codes)
        )
